"""Pull plain text out of uploaded assignment briefs, page by page."""

import re
from dataclasses import dataclass
from pathlib import Path

MAX_PAGES = 200
MAX_CHARS = 300_000
VIRTUAL_PAGE_CHARS = 3_000
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"


class ExtractionError(ValueError):
    """The file cannot be turned into text (scanned image, corrupt, unsupported)."""

    def __init__(self, message: str, reason: str) -> None:
        super().__init__(message)
        self.reason = reason


@dataclass(frozen=True)
class PageText:
    number: int
    text: str


def extract_pages(path: Path, content_type: str) -> list[PageText]:
    if content_type == PDF:
        pages = _pdf_pages(path)
    elif content_type == DOCX:
        pages = _virtual_pages(_docx_text(path))
    elif content_type == TXT:
        pages = _virtual_pages(path.read_text(encoding="utf-8", errors="replace"))
    else:
        raise ExtractionError("Unsupported file type", "unsupported")
    pages = [page for page in pages if page.text.strip()]
    if not pages:
        raise ExtractionError("The file contains no extractable text", "no_text")
    return _limit(pages)


def _limit(pages: list[PageText]) -> list[PageText]:
    kept: list[PageText] = []
    total = 0
    for page in pages[:MAX_PAGES]:
        remaining = MAX_CHARS - total
        if remaining <= 0:
            break
        text = page.text[:remaining]
        kept.append(PageText(page.number, text))
        total += len(text)
    return kept


def _pdf_pages(path: Path) -> list[PageText]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is password protected", "failed")
        return [
            PageText(index, _clean(page.extract_text() or ""))
            for index, page in enumerate(reader.pages[:MAX_PAGES], start=1)
        ]
    except (PyPdfError, ValueError, OSError, KeyError) as error:
        raise ExtractionError("The PDF could not be read", "failed") from error


def _docx_text(path: Path) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(str(path))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as error:
        raise ExtractionError("The DOCX could not be read", "failed") from error
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _virtual_pages(text: str) -> list[PageText]:
    """DOCX and TXT have no pages: group paragraphs into stable ~3,000 character sections."""
    chunks: list[str] = []
    current = ""
    for paragraph in re.split(r"\n\s*\n|\n", text.replace("\r", "")):
        if current and len(current) + len(paragraph) > VIRTUAL_PAGE_CHARS:
            chunks.append(current)
            current = ""
        current = f"{current}\n{paragraph}" if current else paragraph
    if current:
        chunks.append(current)
    return [PageText(index, _clean(chunk)) for index, chunk in enumerate(chunks, start=1)]


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text).strip()
