from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from app.documents.extract import DOCX, PDF, TXT, ExtractionError, extract_pages

SE_COURSE = "00000000-0000-0000-0000-0000000000c8"


def make_pdf(path: Path, pages: list[list[str]]) -> None:
    pdf = canvas.Canvas(str(path))
    for lines in pages:
        y = 800
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()


def test_pdf_text_is_returned_page_by_page(tmp_path: Path) -> None:
    path = tmp_path / "brief.pdf"
    make_pdf(path, [["Assignment 3", "R1 Implement JWT login."], ["R2 Secure endpoints."]])
    pages = extract_pages(path, PDF)
    assert [page.number for page in pages] == [1, 2]
    assert "Implement JWT login." in pages[0].text and "Secure endpoints." in pages[1].text


def test_docx_and_txt_are_grouped_into_stable_sections(tmp_path: Path) -> None:
    docx = tmp_path / "brief.docx"
    document = DocxDocument()
    document.add_paragraph("Requirement one")
    document.add_paragraph("Requirement two")
    document.save(str(docx))
    assert "Requirement two" in extract_pages(docx, DOCX)[0].text

    long_text = tmp_path / "long.txt"
    long_text.write_text("\n\n".join(f"Paragraph {i} " + "x" * 500 for i in range(20)))
    pages = extract_pages(long_text, TXT)
    assert len(pages) > 1 and [p.number for p in pages] == list(range(1, len(pages) + 1))


def test_unreadable_files_report_why(tmp_path: Path) -> None:
    blank = tmp_path / "blank.pdf"
    make_pdf(blank, [[]])
    with pytest.raises(ExtractionError) as empty:
        extract_pages(blank, PDF)
    assert empty.value.reason == "no_text"

    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really")
    with pytest.raises(ExtractionError) as corrupt:
        extract_pages(broken, PDF)
    assert corrupt.value.reason == "failed"

    with pytest.raises(ExtractionError) as other:
        extract_pages(blank, "image/png")
    assert other.value.reason == "unsupported"


def upload(client: TestClient, name: str, data: bytes, content_type: str) -> dict:
    response = client.post(
        "/api/v1/assignments/upload",
        data={"course_id": SE_COURSE, "title": f"Upload {name}", "due_at": "2026-11-01T18:00:00Z"},
        files={"file": (name, data, content_type)},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_upload_indexes_the_text_into_page_fragments(
    student: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core.config import settings

    monkeypatch.setattr(settings, "upload_directory", str(tmp_path / "uploads"))
    path = tmp_path / "brief.pdf"
    make_pdf(path, [["R1 Implement JWT login."], ["R2 Secure endpoints."]])
    body = upload(student, "brief.pdf", path.read_bytes(), PDF)
    assert body["extraction"] == {
        "status": "ok",
        "pages": 2,
        "characters": body["extraction"]["characters"],
    }
    assert body["document"]["page_count"] == 2

    knowledge = student.get("/api/v1/knowledge").json()
    ours = [f for f in knowledge["fragments"] if f["document_id"] == body["document"]["id"]]
    assert sorted(f["page"] for f in ours) == [1, 2]
    assert any("Secure endpoints" in f["text"] for f in ours)


def test_upload_still_succeeds_when_no_text_can_be_extracted(
    student: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core.config import settings

    monkeypatch.setattr(settings, "upload_directory", str(tmp_path / "uploads"))
    body = upload(student, "scan.pdf", b"%PDF-1.4 scanned", PDF)
    assert body["extraction"]["status"] == "failed" and body["document"]["page_count"] == 0
